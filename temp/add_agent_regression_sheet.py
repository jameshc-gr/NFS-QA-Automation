from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


WORKBOOK_PATH = Path(__file__).with_name("Rate App - AUG 2026_Full_Regression_Suite.xlsx")
SHEET_NAME = "16. Agent Regression"
HEADERS = [
    "Issue Keys",
    "Module",
    "Category",
    "Test ID",
    "Test Case",
    "Steps",
    "Expected Result",
    "iOS (Pass/Fail)",
    "Android (Pass/Fail)",
    "QA Assignee",
    "Note",
]

CASES = [
    ("MSAM-8475", "Rate Homes", "Find Agent Form", "AGENT-1", "Find Agent form matches the approved design", "Open Rate Homes; select Find Agent; verify field order; enter a ZIP code; verify state is read-only and county options appear.", "The form matches the approved design, ZIP code is last, state is populated read-only, and county options are available."),
    ("MSAM-8225", "Registration", "Email Verification", "AGENT-2", "New account completes email verification without agent-flag errors", "Create a new account in QA; request and enter a new email verification code; select Verify.", "Verification completes without an agent-flag error, error page, duplicate-registration state, or blocked retry."),
    ("MSAM-8512", "Deep Links", "Personal Collections", "AGENT-3", "Personal Collections deep links open in the iOS app", "On iOS PROD, open https://mobileapp.rate.com/personal-collections; repeat with rate://personal-collections.", "Each supported link opens the Personal Collections screen in the app."),
    ("MSAM-8509", "Agent Invite", "Existing User", "AGENT-4", "Existing Rate App user reaches the agent view from an invite", "Use Power VP to send an agent invite to an existing Rate App user; open the SMS link; sign in with that user.", "After authentication, the user remains in the agent experience rather than returning to the borrower view."),
    ("MSAM-8474", "Deep Links", "Agent Authentication", "AGENT-5", "Authenticated agent reaches the requested deep-link destination", "As an agent, open a supported deep link; authenticate when prompted.", "The authenticated user lands on the consumer-side destination represented by the deep link."),
    ("MSAM-8306", "Agent Profile", "Avatar and Logo", "AGENT-6", "Agent imagery renders from the updated AES schema", "Open every Rate App surface that displays an agent avatar or logo; load an agent whose AES record provides images.headshot and images.branding values.", "Agent imagery renders correctly with the current AES image data and no missing-image regressions."),
    ("MSAM-8290", "Agent Invite", "Android Kochava", "AGENT-7", "Kochava agent invite opens registration promptly with prefilled data", "On Android, open an agent-invite Kochava URL; install or open Rate App; observe the first displayed registration screen.", "The app opens directly to Create Account without a delay and prepopulates the invitee name and email."),
    ("MSAM-8273", "Agent Invite", "Android Session", "AGENT-8", "Second Kochava-link use requires a valid session and preserves the correct persona", "While signed in, open an agent-invite Kochava link; sign out; open the link again; inspect Home and Profile.", "The app does not silently reuse a session, prompts for authentication when required, and never displays agent tabs or an empty profile in a borrower view."),
    ("MSAM-8272", "Agent Invite", "Android Persona", "AGENT-9", "New agent-invite registration shows the correct first-time profile view", "Open an agent-invite Kochava link; complete registration; open Profile.", "The first profile view matches the expected client or borrower experience and does not incorrectly offer the agent view."),
    ("MSAM-8256", "Login", "iOS Session Storage", "AGENT-10", "Login tab switching does not loop validation checks", "On iOS, open Login; switch repeatedly between login tabs; inspect behavior and logs; repeat the Remember Me, agent invitation, and loan officer invitation paths.", "Tab switches remain stable with no endless validation cycle or excessive log entries, and the specified paths continue to work."),
    ("MSAM-8232", "Agent Referral", "iOS Signup", "AGENT-11", "Agent-referral signup opens the expected post-signup view", "On iOS, open an agent-referral Kochava link; complete account creation.", "The newly registered user reaches the expected post-signup view and is not routed to an incorrect borrower or client view."),
    ("MSAM-8076", "Rate Homes", "Post-Submission Entry Point", "AGENT-12", "Rate Homes entry point is removed after requesting an agent", "Submit the Rate Homes form and request an agent; return to the Home tab.", "The Rate Homes entry point is no longer displayed after the request is complete."),
    ("MSAM-8075", "Rate Homes", "Lead Submission", "AGENT-13", "Rate Homes lead submission succeeds", "Complete the Rate Homes form with valid required data and submit the lead.", "The lead is submitted successfully and the app confirms completion without an error."),
]


def main() -> None:
    workbook = load_workbook(WORKBOOK_PATH)
    if "14. Agent Regression" in workbook.sheetnames:
        del workbook["14. Agent Regression"]
    if SHEET_NAME in workbook.sheetnames:
        del workbook[SHEET_NAME]

    source_sheet = workbook["1. Auth & Session"]
    sheet = workbook.create_sheet(SHEET_NAME)
    sheet.sheet_view.showGridLines = source_sheet.sheet_view.showGridLines
    sheet.freeze_panes = "A2"

    for column_index, header in enumerate(HEADERS, start=1):
        source_cell = source_sheet.cell(1, column_index)
        target_cell = sheet.cell(1, column_index, header)
        target_cell.font = copy(source_cell.font) if source_cell.font else Font(bold=True)
        target_cell.fill = copy(source_cell.fill) if source_cell.fill else PatternFill("solid", fgColor="FFFFFF")
        target_cell.border = copy(source_cell.border)
        target_cell.alignment = copy(source_cell.alignment)
        target_cell.number_format = source_cell.number_format

    for row_index, case in enumerate(CASES, start=2):
        issue_key, module, category, test_id, test_case, steps, expected_result = case
        values = [issue_key, module, category, test_id, test_case, steps, expected_result, None, None, None, None]
        for column_index, value in enumerate(values, start=1):
            source_cell = source_sheet.cell(2, min(column_index, source_sheet.max_column))
            target_cell = sheet.cell(row_index, column_index, value)
            target_cell.font = copy(source_cell.font)
            target_cell.fill = copy(source_cell.fill)
            target_cell.border = copy(source_cell.border)
            target_cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = [20, 24, 26, 14, 48, 70, 62, 18, 20, 20, 28]
    for column_index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    for row_index in range(2, len(CASES) + 2):
        sheet.row_dimensions[row_index].height = 60

    workbook.save(WORKBOOK_PATH)


if __name__ == "__main__":
    main()