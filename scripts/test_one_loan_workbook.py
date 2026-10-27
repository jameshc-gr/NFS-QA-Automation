import importlib.util
import tempfile
import unittest
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

spec = importlib.util.spec_from_file_location(
    "one_loan_writer", Path(__file__).with_name("update-one-loan-workbook.py")
)
writer = importlib.util.module_from_spec(spec)
spec.loader.exec_module(writer)


class WorkbookExportTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.path = Path(self.directory.name) / "loans.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"
        sheet.append([None, "Field Description", "Name", "Custom", "Missing", "Duplicate"])
        sheet.append(["Loan #", "Encomass Field ID", 4000, "CX.TEST FIELD", 4001, 4000])
        sheet.append([None, "Description", "Existing description"])
        sheet.append(["OLDDEV", "Values", "Original"])
        sheet["C4"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")
        workbook.create_sheet("Sheet2")["A1"] = "untouched"
        workbook.save(self.path)

    def loan(self, number="NEWDEV", **fields):
        return {"loanNumber": number, "fields": {"305": number, **fields}}

    def test_append_update_alias_duplicates_and_preservation(self):
        report = writer.update_workbook(self.path, [
            self.loan(**{"4000": "New", "CX.TEST.FIELD": "0"}),
        ])
        self.assertEqual(report["loans"][0]["populated"], 3)
        self.assertEqual(report["loans"][0]["missingFieldIds"], ["4001"])
        workbook = load_workbook(self.path)
        self.assertEqual(workbook["Sheet1"]["C4"].value, "Original")
        self.assertEqual(workbook["Sheet1"]["C3"].value, "Existing description")
        self.assertEqual(workbook["Sheet2"]["A1"].value, "untouched")
        self.assertEqual(copy(workbook["Sheet1"]["C5"].fill), copy(workbook["Sheet1"]["C4"].fill))
        writer.update_workbook(self.path, [self.loan(**{"4000": "Updated"})])
        sheet = load_workbook(self.path)["Sheet1"]
        self.assertEqual(sheet.max_row, 5)
        self.assertEqual(sheet["C5"].value, "Updated")
        self.assertIsNone(sheet["D5"].value)
        self.assertEqual(sheet["F5"].value, "Updated")

    def test_multiple_loans_and_formula_like_text(self):
        writer.update_workbook(self.path, [
            self.loan("FIRST", **{"4000": "=1+1"}),
            self.loan("SECOND", **{"4000": "000123"}),
        ])
        sheet = load_workbook(self.path)["Sheet1"]
        self.assertEqual(sheet["C5"].value, "=1+1")
        self.assertEqual(sheet["C5"].data_type, "s")
        self.assertEqual(sheet["C6"].value, "000123")

    def test_invalid_identity_does_not_change_file(self):
        before = self.path.read_bytes()
        with self.assertRaisesRegex(ValueError, "identity"):
            writer.update_workbook(self.path, [
                {"loanNumber": "NEWDEV", "fields": {"305": "WRONG"}},
            ])
        self.assertEqual(self.path.read_bytes(), before)

    def test_ambiguous_alias_does_not_change_file(self):
        before = self.path.read_bytes()
        with self.assertRaisesRegex(ValueError, "Ambiguous"):
            writer.update_workbook(self.path, [
                self.loan(**{"CX.TEST.FIELD": "a", "CX.TEST-FIELD": "b"}),
            ])
        self.assertEqual(self.path.read_bytes(), before)

    def test_duplicate_workbook_rows_are_rejected(self):
        workbook = load_workbook(self.path)
        workbook["Sheet1"]["A5"] = "OLDDEV"
        workbook.save(self.path)
        with self.assertRaisesRegex(ValueError, "Duplicate workbook rows"):
            writer.update_workbook(self.path, [self.loan("OLDDEV")])


if __name__ == "__main__":
    unittest.main()
