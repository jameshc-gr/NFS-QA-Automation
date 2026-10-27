import json
import os
import re
import sys
import tempfile
from collections import defaultdict
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles.cell_style import StyleArray


def load_template(path):
    workbook = load_workbook(path)
    sheet = workbook["Sheet1"]
    if sheet["A2"].value != "Loan #" or sheet["B3"].value != "Description":
        raise ValueError("Expected the existing template: loan numbers in A, IDs in row 2, descriptions in row 3")
    headers = [
        (column, str(sheet.cell(2, column).value).strip())
        for column in range(3, sheet.max_column + 1)
        if sheet.cell(2, column).value is not None
    ]
    if not headers:
        raise ValueError("Sheet1 has no field IDs in row 2")
    return workbook, sheet, headers


def update_workbook(path, loans):
    workbook, sheet, headers = load_template(path)
    before = {
        (tab.title, cell.coordinate): (cell.value, tuple(cell._style or StyleArray()))
        for tab in workbook
        for row in tab.iter_rows()
        for cell in row
    }
    expected = {}
    reports = []
    seen = set()
    normalize = lambda value: re.sub(r"[^a-z0-9]", "", str(value).lower())
    for loan in loans:
        number = loan["loanNumber"]
        if number in seen or not re.fullmatch(r"[A-Za-z0-9_-]+", number):
            raise ValueError(f"Duplicate or invalid loan number: {number}")
        seen.add(number)
        fields = loan["fields"]
        if str(fields.get("305", "")).strip() != number:
            raise ValueError(f"Loan identity mismatch: {number}")
        rows = [
            row for row in range(4, sheet.max_row + 1)
            if str(sheet.cell(row, 1).value).strip() == number
        ]
        if len(rows) > 1:
            raise ValueError(f"Duplicate workbook rows for {number}")
        row = rows[0] if rows else sheet.max_row + 1
        if not rows:
            for column in range(1, sheet.max_column + 1):
                sheet.cell(row, column)._style = copy(sheet.cell(row - 1, column)._style)
            sheet.cell(row, 1).value = number
        aliases = defaultdict(list)
        for key, value in fields.items():
            aliases[normalize(key)].append(value)
        missing = []
        populated = 0
        for column, key in headers:
            if key in fields:
                value = fields[key]
            else:
                candidates = aliases[normalize(key)]
                if len(candidates) > 1:
                    raise ValueError(f"Ambiguous field alias: {key}")
                value = candidates[0] if candidates else None
                if not candidates:
                    missing.append(key)
            if value is not None and not isinstance(value, (str, int, float, bool)):
                raise ValueError(f"Unsupported field value for {key}")
            value = None if value == "" else value
            cell = sheet.cell(row, column)
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"
            expected[(sheet.title, cell.coordinate)] = value
            populated += value is not None
        reports.append({
            "loanNumber": number, "row": row, "populated": populated,
            "missingColumns": len(missing), "missingFieldIds": sorted(set(missing)),
        })
    if not reports:
        raise ValueError("No loan data supplied")
    handle, temporary = tempfile.mkstemp(suffix=".xlsx", dir=path.parent)
    os.close(handle)
    try:
        workbook.save(temporary)
        verified = load_workbook(temporary)
        for tab in verified:
            for cells in tab.iter_rows():
                for cell in cells:
                    key = (tab.title, cell.coordinate)
                    if key in expected and cell.value != expected[key]:
                        raise ValueError(f"Saved value mismatch: {key}")
                    if key in before:
                        value, style = before[key]
                        if key not in expected and cell.value != value:
                            raise ValueError(f"Unrelated cell changed: {key}")
                        if tuple(cell._style or StyleArray()) != style:
                            raise ValueError(f"Formatting changed: {key}")
        for report in reports:
            if verified["Sheet1"].cell(report["row"], 1).value.strip() != report["loanNumber"]:
                raise ValueError("Saved loan number mismatch")
        os.chmod(temporary, path.stat().st_mode & 0o777)
        os.replace(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)
    return {"loans": reports, "verified": True}


if __name__ == "__main__":
    target = Path(sys.argv[1]).resolve()
    if len(sys.argv) > 2 and sys.argv[2] == "--check":
        load_template(target)
    else:
        print(json.dumps(update_workbook(target, json.load(sys.stdin))))
